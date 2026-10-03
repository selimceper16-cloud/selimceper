"""Paketlenmiş uygulamanın (PyInstaller exe) tüm bağımlılıklarının çalıştığını doğrular.

    InovasyonZekasi.exe --selftest
    python main.py --selftest

Ollama veya internet GEREKTİRMEZ: LLM ve embedding yerine sahte (deterministik) modeller
kullanılır; ChromaDB geçici bir klasörde gerçekten oluşturulup sorgulanır. Sonuç hem
çıkış koduyla (0 = başarılı) hem de bir metin dosyasına yazılarak bildirilir; konsolsuz
exe'de stdout olmadığı için CI bu dosyayı okur.
"""

from __future__ import annotations

import os
import shutil
import sys
import tempfile
import traceback
from pathlib import Path

import config


def _check_gui() -> None:
    import customtkinter as ctk

    root = ctk.CTk()
    root.withdraw()
    ctk.CTkTextbox(root).insert("1.0", "test")
    root.update()
    root.destroy()


def _check_readers(tmp: Path) -> None:
    import docx
    from PyPDF2 import PdfWriter

    import database_builder

    d = docx.Document()
    d.add_paragraph("Akıllı sera projesi: nem sensörü ile otomatik havalandırma. " * 5)
    d.save(tmp / "ornek.docx")
    writer = PdfWriter()
    writer.add_blank_page(100, 100)
    with open(tmp / "bos.pdf", "wb") as f:
        writer.write(f)
    assert "Akıllı sera" in database_builder.read_docx(tmp / "ornek.docx")
    database_builder.read_pdf(tmp / "bos.pdf")


def _check_rag_and_chain(tmp: Path) -> None:
    from langchain_chroma import Chroma
    from langchain_core.documents import Document
    from langchain_core.embeddings import DeterministicFakeEmbedding
    from langchain_core.language_models.fake_chat_models import GenericFakeChatModel

    import ai_engine

    config.CHROMA_DIR = tmp / "chroma_db"
    store = Chroma(
        collection_name=config.CHROMA_COLLECTION,
        embedding_function=DeterministicFakeEmbedding(size=32),
        persist_directory=str(config.CHROMA_DIR),
        client_settings=config.chroma_client_settings(),
    )
    store.add_documents(
        [Document(page_content="Toprak nem sensörlü otomatik sulama projesi.",
                  metadata={"baslik": "otomatik sulama", "kategori": "Tarım"})],
        ids=["t1"],
    )
    engine = ai_engine.InnovationEngine()
    engine._store = store
    engine._llm = GenericFakeChatModel(messages=iter(["Proje Adı: Selftest"]))
    out = "".join(engine.generate_idea("Lise", "Tarım", use_web=False))
    assert "Proje Adı: Selftest" in out and "otomatik sulama" in out, out


def _check_search_lib() -> None:
    import search_module

    assert search_module.DDGS is not None, "ddgs paketi pakete dahil edilmemiş"
    with search_module.DDGS():
        pass


def _check_live_ollama() -> None:
    """Gerçek (veya taklit) Ollama sunucusuyla embedding + sohbet akışını dener."""
    from langchain_ollama import ChatOllama, OllamaEmbeddings

    import ai_engine

    engine = ai_engine.InnovationEngine()
    for name in engine.missing_models() or []:
        engine.pull_model(name)
    ok, msg = engine.check_ollama()
    assert ok, msg
    emb = OllamaEmbeddings(model=config.EMBED_MODEL, base_url=config.OLLAMA_BASE_URL)
    assert len(emb.embed_query("test")) > 0
    llm = ChatOllama(model=config.LLM_MODEL, base_url=config.OLLAMA_BASE_URL, num_predict=16)
    assert "".join(c.content for c in llm.stream("Merhaba de.")).strip()


def run_selftest() -> int:
    out_path = Path(os.getenv("INOVASYON_SELFTEST_OUT", config.DATA_DIR / "selftest_sonucu.txt"))
    tmp = Path(tempfile.mkdtemp(prefix="inovasyon_selftest_"))
    checks = [
        ("Arayüz (CustomTkinter/Tk)", _check_gui),
        ("Dosya okuyucular (PDF/DOCX)", lambda: _check_readers(tmp)),
        ("ChromaDB + LangChain akışı", lambda: _check_rag_and_chain(tmp)),
        ("Web arama kütüphanesi (ddgs)", _check_search_lib),
    ]
    if os.getenv("INOVASYON_SELFTEST_LIVE") == "1":
        checks.append((f"Canlı Ollama ({config.OLLAMA_BASE_URL})", _check_live_ollama))
    lines = [f"{config.APP_TITLE} v{config.APP_VERSION} selftest (frozen={config.FROZEN})"]
    failed = 0
    for name, fn in checks:
        try:
            fn()
            lines.append(f"OK    {name}")
        except Exception:  # noqa: BLE001
            failed += 1
            lines.append(f"FAIL  {name}\n{traceback.format_exc()}")
    lines.append("SONUC: " + ("BASARILI" if not failed else f"{failed} HATA"))
    report = "\n".join(lines) + "\n"

    shutil.rmtree(tmp, ignore_errors=True)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(report, encoding="utf-8")
    try:
        sys.stdout.write(report)
        sys.stdout.flush()
    except Exception:  # noqa: BLE001 - konsolsuz exe'de stdout bir dosya olabilir
        pass
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(run_selftest())
