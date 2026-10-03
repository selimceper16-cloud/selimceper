"""Geçmiş proje raporlarını yerel ChromaDB'ye aktaran bağımsız script.

Kullanım:
    python database_builder.py                 # varsayılan: ./gecmis_projeler
    python database_builder.py -d raporlarim   # farklı klasör
    python database_builder.py --reset         # veritabanını sıfırdan kur

Klasör yapısı (alt klasörler opsiyoneldir, varsa kategori etiketi olarak kullanılır):
    gecmis_projeler/
        Tarım/akilli_sera_raporu.pdf
        Sağlık/proje2.docx
        genel_proje.txt

Embedding'ler Ollama üzerinden yerelde üretilir (varsayılan: nomic-embed-text).
Önce:  ollama pull nomic-embed-text
"""

from __future__ import annotations

import argparse
import hashlib
import shutil
import sys
from pathlib import Path
from typing import Callable

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

import config

SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".md", ".docx"}
IGNORED_FILES = {"beni_oku.txt", "readme.txt", "readme.md"}  # klasör açıklamaları proje değildir
LogFn = Callable[[str], None]


# ---------------------------------------------------------------------- #
# Dosya okuyucular
# ---------------------------------------------------------------------- #
def read_pdf(path: Path) -> str:
    from PyPDF2 import PdfReader

    reader = PdfReader(str(path))
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:  # noqa: BLE001 - bozuk sayfa tüm dosyayı düşürmesin
            continue
    return "\n".join(pages)


def read_docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def read_text(path: Path) -> str:
    for encoding in ("utf-8", "utf-8-sig", "cp1254", "latin-1"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return ""


READERS = {".pdf": read_pdf, ".docx": read_docx, ".txt": read_text, ".md": read_text}


def file_hash(path: Path) -> str:
    h = hashlib.sha1()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1 << 16), b""):
            h.update(block)
    return h.hexdigest()


def infer_category(path: Path, root: Path) -> str:
    """gecmis_projeler/<Kategori>/dosya.pdf -> 'Kategori'; kök dizindekiler 'Genel'."""
    rel = path.relative_to(root)
    return rel.parts[0] if len(rel.parts) > 1 else "Genel"


# ---------------------------------------------------------------------- #
# Ana akış
# ---------------------------------------------------------------------- #
def load_documents(root: Path, log: LogFn = print) -> list[Document]:
    docs: list[Document] = []
    files = sorted(p for p in root.rglob("*")
                   if p.suffix.lower() in SUPPORTED_EXTENSIONS and p.name.lower() not in IGNORED_FILES)
    if not files:
        log(f"[!] '{root}' içinde desteklenen dosya bulunamadı ({', '.join(SUPPORTED_EXTENSIONS)}).")
        return docs

    for path in files:
        try:
            text = READERS[path.suffix.lower()](path)
        except Exception as exc:  # noqa: BLE001
            log(f"  [x] {path.name}: okunamadı ({exc})")
            continue
        text = " ".join(text.split())  # PDF'lerden gelen fazla boşluk/satır sonlarını temizle
        if len(text) < 50:
            log(f"  [-] {path.name}: metin çok kısa veya taranmış (OCR'sız) PDF, atlandı.")
            continue
        docs.append(Document(
            page_content=text,
            metadata={
                "source": str(path.relative_to(root)),
                "baslik": path.stem.replace("_", " "),
                "kategori": infer_category(path, root),
                "file_hash": file_hash(path),
            },
        ))
        log(f"  [+] {path.relative_to(root)}  ({len(text):,} karakter)")
    return docs


def build(root: Path, reset: bool = False, log: LogFn = print) -> int:
    """Yeni dosyaları veritabanına ekler; eklenen parça sayısını döner."""
    from langchain_chroma import Chroma
    from langchain_ollama import OllamaEmbeddings

    if not root.exists():
        root.mkdir(parents=True)
        log(f"[i] '{root}' klasörü oluşturuldu. Geçmiş proje dosyalarını içine koyup tekrar çalıştırın.")
        return 0

    if reset and config.CHROMA_DIR.exists():
        shutil.rmtree(config.CHROMA_DIR)
        log(f"[i] Eski veritabanı silindi: {config.CHROMA_DIR}")

    log(f"[1/3] Dosyalar okunuyor: {root}")
    docs = load_documents(root, log)
    if not docs:
        return 0

    embeddings = OllamaEmbeddings(model=config.EMBED_MODEL, base_url=config.OLLAMA_BASE_URL)
    store = Chroma(
        collection_name=config.CHROMA_COLLECTION,
        embedding_function=embeddings,
        persist_directory=str(config.CHROMA_DIR),
        client_settings=config.chroma_client_settings(),
    )

    # Daha önce eklenmiş dosyaları atla (dosya içeriği hash'i ile)
    existing = store.get(include=["metadatas"])
    known_hashes = {m.get("file_hash") for m in existing.get("metadatas", []) if m}
    new_docs = [d for d in docs if d.metadata["file_hash"] not in known_hashes]
    skipped = len(docs) - len(new_docs)
    if skipped:
        log(f"[i] {skipped} dosya zaten veritabanında, atlandı.")
    if not new_docs:
        log("[✓] Eklenecek yeni dosya yok. Veritabanı güncel.")
        return 0

    log(f"[2/3] {len(new_docs)} dosya parçalanıyor (chunk={config.CHUNK_SIZE}, overlap={config.CHUNK_OVERLAP})")
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=config.CHUNK_SIZE,
        chunk_overlap=config.CHUNK_OVERLAP,
        separators=["\n\n", "\n", ". ", "; ", ", ", " ", ""],
    )
    chunks = splitter.split_documents(new_docs)
    for idx, chunk in enumerate(chunks):
        chunk.metadata["chunk"] = idx
    ids = [f"{c.metadata['file_hash']}-{i}" for i, c in enumerate(chunks)]

    log(f"[3/3] {len(chunks)} parça için embedding üretilip ChromaDB'ye yazılıyor ({config.EMBED_MODEL})…")
    batch = 64
    for start in range(0, len(chunks), batch):
        store.add_documents(chunks[start:start + batch], ids=ids[start:start + batch])
        log(f"      {min(start + batch, len(chunks))}/{len(chunks)}")

    total = len(store.get(include=[])["ids"])
    log(f"[✓] Tamamlandı. Veritabanında toplam {total} parça var → {config.CHROMA_DIR}")
    return len(chunks)


def main() -> int:
    parser = argparse.ArgumentParser(description="Geçmiş proje raporlarını ChromaDB'ye aktarır.")
    parser.add_argument("-d", "--dir", type=Path, default=config.PROJECTS_DIR,
                        help=f"Rapor klasörü (varsayılan: {config.PROJECTS_DIR.name})")
    parser.add_argument("--reset", action="store_true", help="Mevcut veritabanını silip yeniden kur")
    args = parser.parse_args()
    try:
        build(args.dir.resolve(), reset=args.reset)
    except Exception as exc:  # noqa: BLE001
        from ai_engine import friendly_error

        print(f"[x] {friendly_error(exc)}")
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
